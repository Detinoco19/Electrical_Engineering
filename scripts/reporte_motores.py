
import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def exportar_reporte(
    motores,
    resumen,
    historial,
    calidad,
    archivo_salida,
):
    archivo_salida.parent.mkdir(parents=True, exist_ok=True)

    try:
        with pd.ExcelWriter(archivo_salida, engine="openpyxl") as writer:
            motores.to_excel(writer, sheet_name="MOTORES", index=False)
            resumen.to_excel(writer, sheet_name="RESUMEN", index=False)
            historial.to_excel(writer, sheet_name="HISTORIAL", index=False)
            calidad.to_excel(writer, sheet_name="CALIDAD_DATOS", index=False)

        formatear_excel(archivo_salida)

    except PermissionError as error:
        raise PermissionError(
            "\nNo se puede escribir el archivo:\n"
            f"{archivo_salida}\n\n"
            "Cierra motores_calculados.xlsx y vuelve a ejecutar."
        ) from error


def _ajustar_anchos(ws):
    for numero_columna in range(1, ws.max_column + 1):
        letra = get_column_letter(numero_columna)
        maximo = 0

        for fila in range(1, ws.max_row + 1):
            valor = ws.cell(row=fila, column=numero_columna).value
            if valor is not None:
                maximo = max(maximo, len(str(valor)))

        ws.column_dimensions[letra].width = min(max(maximo + 3, 12), 35)


def _formatear_hoja_datos(
    ws,
    encabezado_fill,
    encabezado_font,
    borde,
    verde,
    amarillo,
    rojo,
    azul,
    gris,
):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 45

    for celda in ws[1]:
        celda.fill = encabezado_fill
        celda.font = encabezado_font
        celda.border = borde
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.border = borde
            celda.alignment = Alignment(vertical="center")

    encabezados = {celda.value: celda.column for celda in ws[1] if celda.value}

    for nombre in ["FECHA", "FECHA_FABRICACION"]:
        if nombre in encabezados:
            columna = encabezados[nombre]
            for fila in range(2, ws.max_row + 1):
                ws.cell(row=fila, column=columna).number_format = "dd/mm/yyyy"

    columnas_numericas = [
        "KW_PLACA",
        "POTENCIA_CV_PLACA",
        "HP",
        "VOLTAJE_V",
        "CORRIENTE_PLACA_A",
        "VOLTAJE_PLACA_1_V",
        "CORRIENTE_PLACA_1_A",
        "VOLTAJE_PLACA_2_V",
        "CORRIENTE_PLACA_2_A",
        "FRECUENCIA_HZ",
        "FP",
        "EFICIENCIA",
        "RPM_MOTOR",
        "FACTOR_SERVICIO",
        "ELEVACION_TEMP_K",
        "TEMP_AMBIENTE_MIN_C",
        "TEMP_AMBIENTE_MAX_C",
        "ALTITUD_MAX_M",
        "IP_IN",
        "CORRIENTE_FS_1_A",
        "CORRIENTE_FS_2_A",
        "RPM_SALIDA",
        "RELACION_REDUCCION",
        "PAR_SALIDA_NM",
        "VOLUMEN_ACEITE_L",
        "VOLTAJE_FRENO_V",
        "PAR_FRENO_NM",
        "PESO_KG",
        "CORRIENTE_L1_A",
        "CORRIENTE_L2_A",
        "CORRIENTE_L3_A",
        "VOLTAJE_L1_L2_V",
        "VOLTAJE_L2_L3_V",
        "VOLTAJE_L3_L1_V",
        "TEMPERATURA_C",
        "AISLAMIENTO_MOHM",
        "KW_CALCULADO_HP",
        "CORRIENTE_CALCULADA_A",
        "DESVIACION_CALC_PLACA_PCT",
        "CORRIENTE_PROMEDIO_A",
        "DESBALANCE_CORRIENTE_PCT",
        "VOLTAJE_PROMEDIO_V",
        "DESBALANCE_VOLTAJE_PCT",
        "CARGA_CORRIENTE_PCT",
    ]

    for nombre in columnas_numericas:
        if nombre not in encabezados:
            continue
        columna = encabezados[nombre]
        for fila in range(2, ws.max_row + 1):
            ws.cell(row=fila, column=columna).number_format = "0.00"

    _ajustar_anchos(ws)

    anchos_especiales = {
        "DESCRIPCION": 40,
        "OBSERVACIONES": 50,
        "OBSERVACIONES_ACTIVO": 50,
        "OBSERVACIONES_MEDICION": 50,
        "DIAGNOSTICO_GENERAL": 42,
        "MODELO": 30,
    }

    for nombre, ancho in anchos_especiales.items():
        if nombre in encabezados:
            letra = get_column_letter(encabezados[nombre])
            ws.column_dimensions[letra].width = ancho

    for nombre in [
        "ALERTA_CARGA",
        "ALERTA_DESBALANCE_CORRIENTE",
        "ALERTA_DESBALANCE_VOLTAJE",
        "DIAGNOSTICO_GENERAL",
    ]:
        if nombre not in encabezados:
            continue

        columna = encabezados[nombre]

        for fila in range(2, ws.max_row + 1):
            celda = ws.cell(row=fila, column=columna)
            valor = str(celda.value or "")

            if valor == "NORMAL":
                celda.fill = verde
            elif valor == "CARGA ALTA":
                celda.fill = amarillo
            elif valor == "CARGA BAJA":
                celda.fill = azul
            elif valor == "SIN MEDICION":
                celda.fill = gris
            elif valor:
                celda.fill = rojo

            celda.font = Font(bold=True)
            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )


def _formatear_calidad(
    ws,
    encabezado_fill,
    encabezado_font,
    borde,
    verde,
    amarillo,
    rojo,
    azul,
):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30

    for celda in ws[1]:
        celda.fill = encabezado_fill
        celda.font = encabezado_font
        celda.border = borde
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.border = borde
            celda.alignment = Alignment(vertical="center", wrap_text=True)

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 75

    for fila in range(2, ws.max_row + 1):
        celda = ws.cell(row=fila, column=1)
        nivel = str(celda.value or "")

        if nivel == "ADVERTENCIA":
            celda.fill = amarillo
        elif nivel == "INFORMACION":
            celda.fill = azul
        elif nivel == "ERROR":
            celda.fill = rojo
        elif nivel == "OK":
            celda.fill = verde

        celda.font = Font(bold=True)


def formatear_excel(archivo_salida):
    wb = load_workbook(archivo_salida)

    encabezado_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    encabezado_font = Font(color="FFFFFF", bold=True)

    lado = Side(style="thin", color="BFBFBF")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)

    verde = PatternFill(fill_type="solid", fgColor="C6EFCE")
    amarillo = PatternFill(fill_type="solid", fgColor="FFEB9C")
    rojo = PatternFill(fill_type="solid", fgColor="FFC7CE")
    azul = PatternFill(fill_type="solid", fgColor="D9EAF7")
    gris = PatternFill(fill_type="solid", fgColor="E7E6E6")

    for nombre in ["MOTORES", "HISTORIAL"]:
        if nombre in wb.sheetnames:
            _formatear_hoja_datos(
                wb[nombre],
                encabezado_fill,
                encabezado_font,
                borde,
                verde,
                amarillo,
                rojo,
                azul,
                gris,
            )

    if "RESUMEN" in wb.sheetnames:
        ws = wb["RESUMEN"]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.column_dimensions["A"].width = 45
        ws.column_dimensions["B"].width = 22
        ws.row_dimensions[1].height = 30

        for celda in ws[1]:
            celda.fill = encabezado_fill
            celda.font = encabezado_font
            celda.border = borde
            celda.alignment = Alignment(horizontal="center", vertical="center")

        for fila in ws.iter_rows(min_row=2):
            for celda in fila:
                celda.border = borde

        for fila in range(2, ws.max_row + 1):
            ws[f"B{fila}"].number_format = "0.00"

    if "CALIDAD_DATOS" in wb.sheetnames:
        _formatear_calidad(
            wb["CALIDAD_DATOS"],
            encabezado_fill,
            encabezado_font,
            borde,
            verde,
            amarillo,
            rojo,
            azul,
        )

    wb.save(archivo_salida)
